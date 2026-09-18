"""Writes docs/Metric_definitions.xlsx - one sheet, one row per figure on the
dashboard: what it means, the formula as calculated, the report and the exact
export columns it comes from, what the pipeline does to those columns first,
which rows count, and what to watch out for.

Reads docs/doc_facts.json (run docs/export_doc_facts.py first), so the sheet,
the Word pack and the dashboard always say the same thing.

Run from the project root:
    python docs/export_doc_facts.py && python docs/build_metric_definitions.py
"""
import json
import pathlib

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).resolve().parent.parent
FACTS = ROOT / "docs" / "doc_facts.json"
OUT = ROOT / "docs" / "Metric_definitions.xlsx"

NAVY, BAND, GREY = "FF1F3864", "FFD9E2F3", "FF595959"
_thin = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

HEADERS = [
    ("Metric", 26), ("Where it appears", 20), ("What it means", 40), ("How it is calculated", 40),
    ("Report", 10), ("Columns in the export", 38), ("What the pipeline does to those columns first", 52),
    ("Rows included", 34), ("Counted per", 12), ("Value", 16), ("Watch out for", 46),
]
# The order the sheet runs in, matching the dashboard top to bottom.
SECTION_ORDER = ["Overview", "The funnel", "How much we wrote", "How good the business is", "Pricing",
                 "What kind of book we write", "Productivity (stand-in headcount)", "Underwriters"]


def main():
    facts = json.loads(FACTS.read_text(encoding="utf-8"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Metric definitions"

    note = (f"Definitions for the UW Productivity dashboard. One row per figure on the page, in the order "
            f"the page shows them. \"Columns in the export\" are the real column headings in DSR and RBS, "
            f"so any figure can be traced back to the report it came from. \"What the pipeline does\" is "
            f"everything that happens to those columns between the export and the figure - that is where "
            f"the original reports are changed. Values are for {facts['period']} on the whole book "
            f"(no filters), data as at {facts['as_at']}. The metrics workbook has the same material with "
            f"the decisions and checks alongside it.")
    cell = ws.cell(1, 1, note)
    cell.font = Font(name="Arial", size=9.5, color=GREY)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws.row_dimensions[1].height = 58

    for i, (title, width) in enumerate(HEADERS, 1):
        head = ws.cell(2, i, title)
        head.font = Font(name="Arial", size=10.5, bold=True, color="FFFFFFFF")
        head.fill = PatternFill("solid", fgColor=NAVY)
        head.alignment = Alignment(wrap_text=True, vertical="center")
        head.border = BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 30

    by_section = {}
    for metric in facts["metrics"]:
        by_section.setdefault(metric["section"], []).append(metric)

    row = 3
    for section in SECTION_ORDER + [s for s in by_section if s not in SECTION_ORDER]:
        for metric in by_section.get(section, []):
            reports = " + ".join(sorted({c["report"] for c in metric["columns"]}))
            columns = "\n".join(f"{c['report']}: {' / '.join(c['names'])}" for c in metric["columns"])
            steps = "\n\n".join(
                f"{c['report']} {c['label']}:\n" + "\n".join(f"- {s}" for s in c["steps"])
                for c in metric["columns"])
            watch = "\n".join(f"- {c}" for c in metric["catches"]) or "Nothing beyond the rows above."
            where = section if metric.get("on_the_page", True) else f"{section} (shown as a note)"
            values = [metric["label"] or metric["title"], where, metric["meaning"], metric["formula"],
                      reports, columns, steps, metric["rows"], metric["counted_per"],
                      metric["now"] or "-", watch]
            for i, value in enumerate(values, 1):
                out = ws.cell(row, i, value)
                out.font = Font(name="Arial", size=9.5, bold=i == 1)
                out.alignment = Alignment(wrap_text=True, vertical="top")
                out.border = BORDER
            row += 1

    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(HEADERS))}{row - 1}"
    wb.save(OUT)
    print(f"Wrote {OUT.name}: {row - 3} metrics on one sheet.")


if __name__ == "__main__":
    main()
