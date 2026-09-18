"""Rebuilds tabs 10 ("Metric lineage") and 11 ("Source columns") of the
metrics workbook from metrics/lineage.py - the same source the dashboard's
(i) help reads - so the Excel and the dashboard always say the same thing.

Only those two tabs are replaced; every other tab is left exactly as it is.
If the real extracts are in data/, tab 11 also shows how full each column is.

Run from the project root after changing metrics/lineage.py:
    python docs/update_lineage_tabs.py
"""
import datetime
import pathlib
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from metrics.lineage import COLUMNS, DSR, METRICS, RBS  # noqa: E402

WORKBOOK = ROOT / "docs" / "UW_Dashboard_Metrics_Inventory.xlsx"
LINEAGE_TAB, COLUMNS_TAB = "10. Metric lineage", "11. Source columns"

NAVY, SECTION, GREY_TEXT = "FF1F3864", "FFD9E2F3", "FF595959"
GREEN, AMBER = "FFE2EFDA", "FFFFF2CC"
_thin = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

SECTION_NAMES = {
    "funnel": "THE FUNNEL", "premium": "HOW MUCH WE WROTE", "quality": "HOW GOOD THE BUSINESS IS",
    "pricing": "PRICING", "composition": "WHAT KIND OF BOOK WE WRITE",
    "productivity_stand_in": "PRODUCTIVITY (STAND-IN HEADCOUNT)", "underwriters": "UNDERWRITER TABLE",
}


def _fresh_sheet(wb, title):
    """Replace a tab in place (same position), or add it at the end."""
    index = len(wb.sheetnames)
    if title in wb.sheetnames:
        index = wb.sheetnames.index(title)
        del wb[title]
    return wb.create_sheet(title, index)


def _intro(ws, text, ncols):
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="Arial", size=9.5, color=GREY_TEXT)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.row_dimensions[1].height = 66


def _header(ws, row, titles, widths):
    for i, (title, width) in enumerate(zip(titles, widths), 1):
        cell = ws.cell(row, i, title)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def _section(ws, row, text, ncols):
    for i in range(1, ncols + 1):
        cell = ws.cell(row, i, text if i == 1 else None)
        cell.font = Font(name="Arial", size=10.5, bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=SECTION)
        cell.border = BORDER
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)


def _row(ws, row, values, fills=None):
    for i, value in enumerate(values, 1):
        cell = ws.cell(row, i, value)
        cell.font = Font(name="Arial", size=9.5, bold=i == 1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER
    for col, colour in (fills or {}).items():
        ws.cell(row, col).fill = PatternFill("solid", fgColor=colour)


def _bullets(items):
    return "\n".join(f"• {item}" for item in items)


def _columns_text(metric):
    """Original columns per report, e.g. "RBS: Agency GELR (%)"."""
    lines, seen = [], set()
    for report, shared in metric.columns:
        if (report, shared) not in seen:
            seen.add((report, shared))
            lines.append(f"{report}: " + " (else, where blank) ".join(COLUMNS[shared].raw(report)))
    return "\n".join(lines)


def _changes_text(metric):
    """What the pipeline does to each of those columns."""
    blocks, seen = [], set()
    for report, shared in metric.columns:
        if (report, shared) not in seen:
            seen.add((report, shared))
            column = COLUMNS[shared]
            blocks.append(f"{report} {column.label}:\n" + _bullets(column.steps_for(report)))
    return "\n\n".join(blocks)


def build_lineage_tab(wb, today):
    ws = _fresh_sheet(wb, LINEAGE_TAB)
    _intro(ws, "Every figure on the dashboard, traced to the original DSR / RBS columns it is built from, "
               "what the pipeline does to each of those columns on the way in, which rows count, and the "
               "formula. This is exactly what the (i) next to each figure on the dashboard shows - both are "
               "generated from metrics/lineage.py, so they can't disagree. Rows marked \"Yes\" under "
               "\"Checked\" are rebuilt straight from the raw columns, without any dashboard code, by "
               "tests/test_trace_to_source.py, which fails if the two ever differ. Before any figure is "
               "worked out, your filters narrow the rows - the columns they use are in tab 11. "
               f"Generated {today} by docs/update_lineage_tabs.py - don't edit this tab by hand.", 11)
    titles = ["Figure (as on the dashboard)", "What it tells you", "Report", "Original columns",
              "What the pipeline does to them", "Rows used", "How it's worked out", "Worth knowing",
              "Counted per", "Checked against raw columns?", "Code"]
    _header(ws, 2, titles, [22, 26, 9, 34, 52, 30, 30, 36, 10, 12, 26])
    row, last_section = 3, None
    for key, m in METRICS.items():
        section = key.split(".")[0]
        if section != last_section:
            _section(ws, row, SECTION_NAMES.get(section, section.upper()), len(titles))
            row, last_section = row + 1, section
        checked = "Yes - exact match" if m.traced else "Not yet"
        _row(ws, row, [m.title, m.meaning, " + ".join(m.reports()), _columns_text(m), _changes_text(m),
                       m.rows, m.formula, _bullets(m.catches), m.counted_per, checked, m.code],
             fills={10: GREEN if m.traced else AMBER})
        row += 1
    ws.freeze_panes = "B3"


def _fill_rates():
    """Share of rows with a value, per original column - only if the real extracts are present."""
    if not (ROOT / "data" / "DSR.xlsx").exists():
        return None, None
    from data_sources.excel_source import ExcelSource
    source = ExcelSource(folder=str(ROOT / "data"))
    dsr, rbs = source.get_dsr(), source.get_rbs()
    as_at = dsr["Submission Date"].max()
    rates = {DSR: {c: dsr[c].notna().mean() for c in dsr.columns},
             RBS: {c: rbs[c].notna().mean() for c in rbs.columns}}
    return rates, f"extract as at {as_at:%d %b %Y} (DSR {len(dsr):,} rows, RBS {len(rbs):,} rows)"


def build_columns_tab(wb, today):
    rates, extract = _fill_rates()
    ws = _fresh_sheet(wb, COLUMNS_TAB)
    _intro(ws, "Every column the dashboard reads from the two exports: its original name in each report, "
               "what the pipeline does to it on the way in, and which figures use it. Any column not listed "
               "is not used. If a \"Required\" column is missing from a new extract, the refresh stops with a "
               "message naming it; if an \"Optional\" one is missing, only the figures that use it go blank. "
               + (f"\"Filled\" = share of rows with a value in the {extract}. " if extract else "")
               + f"Generated {today} by docs/update_lineage_tabs.py from metrics/lineage.py - don't edit "
               "this tab by hand.", 9)
    titles = ["Shared name in the code", "What it is", "DSR column", "RBS column",
              "What the pipeline does to it", "Used by (figures)", "If missing", "Filled in DSR",
              "Filled in RBS"]
    _header(ws, 2, titles, [20, 18, 28, 32, 56, 40, 11, 10, 10])
    used_by = {shared: [] for shared in COLUMNS}
    for m in METRICS.values():
        for _, shared in m.columns:
            if m.title not in used_by[shared]:
                used_by[shared].append(m.title)

    def filled(report, names):
        if not rates or not names:
            return "-"
        return " / ".join(f"{rates[report].get(n, 0):.1%}" for n in names)

    for row, (shared, c) in enumerate(COLUMNS.items(), start=3):
        steps = []
        if c.steps:
            steps.append(_bullets(c.steps))
        if c.dsr and c.dsr_steps:
            steps.append("DSR only:\n" + _bullets(c.dsr_steps))
        if c.rbs and c.rbs_steps:
            steps.append("RBS only:\n" + _bullets(c.rbs_steps))
        uses = ", ".join(used_by[shared]) or "Filters only"
        _row(ws, row, [shared, c.label, "\n".join(c.dsr) or "- (not used)", "\n".join(c.rbs) or "- (not used)",
                       "\n\n".join(steps), uses, "Required" if c.required else "Optional",
                       filled(DSR, c.dsr), filled(RBS, c.rbs)],
             fills={7: AMBER if c.required else GREEN})
    ws.freeze_panes = "B3"


def main():
    today = datetime.date.today().strftime("%d %b %Y")
    wb = openpyxl.load_workbook(WORKBOOK)
    build_lineage_tab(wb, today)
    build_columns_tab(wb, today)
    wb.save(WORKBOOK)
    print(f"Rebuilt '{LINEAGE_TAB}' ({len(METRICS)} figures) and '{COLUMNS_TAB}' ({len(COLUMNS)} columns) "
          f"in {WORKBOOK.name}.")


if __name__ == "__main__":
    main()
